"""生成《软件输出自动化》产品需求文档 (.docx)"""
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from pathlib import Path

doc = Document()

# ── 全局样式 ──
style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.5
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

for level in range(1, 4):
    h = doc.styles[f'Heading {level}']
    h.font.name = 'Microsoft YaHei'
    h.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    h.font.color.rgb = RGBColor(0x1a, 0x56, 0xc4)


def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    return table


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 封面
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
for _ in range(6):
    doc.add_paragraph()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('软件输出自动化')
run.font.size = Pt(28)
run.bold = True
run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('产品需求文档 (PRD)')
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.add_paragraph()

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = info.add_run('版本：v1.0\n文档日期：2026-07-21\n编写：产品部')
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

doc.add_page_break()

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 目录页
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
doc.add_heading('目录', level=1)
toc_items = [
    '1. 项目概述',
    '   1.1 背景',
    '   1.2 产品目标',
    '   1.3 核心价值',
    '   1.4 产品边界与限制',
    '2. 用户角色与使用场景',
    '3. 功能需求',
    '   3.1 自动模式',
    '   3.2 手动模式',
    '   3.3 审核面板',
    '   3.4 执行结果',
    '   3.5 编译队列',
    '   3.6 编译构建',
    '   3.7 规则管理',
    '   3.8 网盘上传',
    '   3.9 代码提交',
    '4. 界面设计要求',
    '5. 非功能需求',
    '6. 技术约束',
    '7. 版本规划',
]
for item in toc_items:
    p = doc.add_paragraph(item)
    p.paragraph_format.space_after = Pt(2)

doc.add_page_break()

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 1. 项目概述
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
doc.add_heading('1. 项目概述', level=1)

doc.add_heading('1.1 背景', level=2)
doc.add_paragraph(
    '在 Android TV（CTV）客户定制业务中，工程师需要对客户的软件配置进行大量重复性修改，'
    '包括语言设置、背光电流、杜比开关、白名单管理等十余类配置项。'
    '每个客户的定制需求涉及多个配置文件（build_config.txt、db.ini、ctv_data.xml 等），'
    '手动修改效率低、容易出错。'
)

doc.add_heading('1.2 产品目标', level=2)
doc.add_paragraph(
    '「软件输出自动化」是一款桌面工具软件，目标是将 CTV 客户定制流程从「人工逐项修改」'
    '升级为「规则匹配 + 一键执行 + 自动编译 + 一键上传」的全流程自动化。'
)

doc.add_heading('1.3 核心价值', level=2)
values = [
    ('效率提升', '单个客户定制从 30 分钟缩短到 3 分钟'),
    ('减少出错', '规则引擎自动匹配，避免人工遗漏'),
    ('流程标准化', '修改 → 审核 → 编译 → 上传，全链路可追溯'),
    ('知识沉淀', '修改规则可配置、可复用，不依赖个人经验'),
]
add_table(['价值维度', '说明'], values)

doc.add_heading('1.4 产品边界与限制', level=2)
doc.add_paragraph(
    '本软件针对已建立标准化工作流的 CTV 客户定制场景设计，适用于批量、重复性的配置修改任务。'
    '以下场景超出本软件的能力范围，需要人工处理：'
)
limits = [
    ('涉及图片/资源文件的修改', '如修改 logo、开机画面、壁纸等需要设计素材的需求'),
    ('需求以图片形式提供', '如客户通过截图、照片传递需求，系统无法解析图片内容'),
    ('首次软件需求（全新项目）', '首次定制涉及整体架构搭建、目录创建、基础配置，需人工完成'),
    ('遥控器相关修改', '遥控器按键映射、配对协议等硬件层配置，需专用工具和手动调试'),
    ('非标准化的特殊需求', '如内核驱动修改、硬件适配、第三方 SDK 集成等超出配置文件范围的需求'),
]
add_table(['不可用场景', '说明'], limits)
doc.add_paragraph(
    '总结：本软件执行的是客户工作流已顺畅、修改规则已沉淀的标准场景。'
    '对于非标准化需求，建议先由工程师人工完成首次配置，后续同类需求再通过本软件自动化。'
)

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 2. 用户角色
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
doc.add_heading('2. 用户角色与使用场景', level=1)

doc.add_heading('2.1 用户角色', level=2)
roles = [
    ('FAE 工程师', '主要用户', '输入客户需求、执行修改、编译、上传软件'),
    ('项目经理', '审核者', '审核 AI 分析结果，确认修改是否正确'),
    ('测试工程师', '次要用户', '使用手动模式逐项调整配置进行测试'),
]
add_table(['角色', '定位', '职责'], roles)

doc.add_heading('2.2 典型使用场景', level=2)

doc.add_heading('场景一：快速定制（自动模式）', level=3)
steps = [
    'FAE 收到客户需求："默认语言英语英国，背光 680mA，关蓝屏"',
    '在自动模式输入需求文本，点击「分析」',
    '系统自动匹配规则，识别 5 条修改项，提交审核',
    '审核通过后，系统自动修改远程服务器上的配置文件',
    '修改完成，加入编译队列',
    '编译成功后，上传到公司网盘，生成分享链接',
]
for i, s in enumerate(steps, 1):
    doc.add_paragraph(f'{i}. {s}')

doc.add_heading('场景二：精细调整（手动模式）', level=3)
doc.add_paragraph(
    '测试工程师需要单独调整白平衡参数，在手动模式下切换到 db.ini 页面，'
    '直接修改 R/G/B 值，点击「执行修改」即可。'
)

doc.add_heading('场景三：批量编译（编译队列）', level=3)
doc.add_paragraph(
    '同时有 10 个客户的软件需要编译。将所有客户加入编译队列，'
    '在需要 clean 的位置插入 clean 命令卡片，点击「全部开始编译」，'
    '系统自动按顺序执行。'
)

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 3. 功能需求
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
doc.add_heading('3. 功能需求', level=1)

# ── 3.1 自动模式 ──
doc.add_heading('3.1 自动模式', level=2)
doc.add_paragraph('用户输入自然语言需求，系统自动识别修改项并执行。')

doc.add_heading('功能描述', level=3)
features = [
    ('需求输入', '支持手动输入、导入文件两种方式', '高'),
    ('规则匹配', '基于关键词匹配内置规则，识别客户需求中的修改项', '高'),
    ('上下文感知', '根据当前项目/板型/区域/客户目录提供精确分析', '高'),
    ('模糊匹配', '"菜单显示时间5秒" 中间有空格也能识别', '中'),
    ('否定识别', '"不显示分辨率" 识别为隐藏，而非显示', '高'),
    ('AI 智能分析（规划中）', '接入 LLM 实现自然语言理解，提升识别准确率', '远期'),
]
add_table(['功能', '说明', '优先级'], features)

doc.add_heading('输入/输出', level=3)
add_table(['项目', '说明'], [
    ('输入', '客户需求文本（自然语言）'),
    ('输出', 'AnalysisResult：客户名、平台、修改项列表、涉及文件'),
])

# ── 3.2 手动模式 ──
doc.add_heading('3.2 手动模式', level=2)
doc.add_paragraph('提供可视化的配置编辑界面，用户可逐项查看和修改远程服务器上的配置。')

doc.add_heading('页面结构', level=3)
pages = [
    ('build_config.txt', '功能开关 + 参数', '杜比、投屏、Eshare、背光电流'),
    ('db.ini', '蓝屏 + 高级参数', '蓝屏开关、白平衡、NLA、色温'),
    ('ctvbuild.prop', '上电/开机模式', '上电模式（待机/开机/记忆）、开机模式'),
    ('ctv_data.xml', '桌面/菜单/语言', '开机桌面、菜单时间、默认语言/国家'),
    ('ctvsetting.xml', '菜单项显示', '各菜单项的显示/隐藏'),
    ('whiteList.conf', '白名单管理', '查看、添加、移除白名单包名'),
    ('build_ctv_app.txt', '预装应用', 'ESharePlus 预装/取消'),
]
add_table(['配置文件', '功能分类', '主要配置项'], pages)

doc.add_heading('交互规范', level=3)
specs = [
    ('二选一配置', '开关按钮组（打开/关闭），显示当前值'),
    ('数值配置', '输入框，显示当前值，可直接修改'),
    ('多选配置', '下拉框（如语言列表、国家列表）'),
    ('预览', '执行前可预览将要修改的内容'),
    ('实时读取', '切换页面时自动从服务器读取最新配置值'),
]
add_table(['配置类型', '交互方式'], specs)

# ── 3.3 审核面板 ──
doc.add_heading('3.3 审核面板', level=2)
doc.add_paragraph('AI 分析完成后，展示分析结果供人工审核确认。')

review_items = [
    ('基本信息', '项目、版型、区域、客户、操作模式'),
    ('修改项列表', '每条修改的规则名称、目标值、涉及文件'),
    ('校验警告', '如目标目录已存在、配置值异常等'),
    ('操作按钮', '通过 → 执行修改 / 拒绝 → 返回（需填写理由）'),
]
add_table(['区域', '内容'], review_items)

# ── 3.4 执行结果 ──
doc.add_heading('3.4 执行结果', level=2)
doc.add_paragraph('修改执行完成后，展示详细的修改结果。')

result_items = [
    ('修改概要', '所有修改项的列表'),
    ('文件变更', '被修改的文件列表'),
    ('修改详情', '每个文件的 diff 对比（删除行/新增行）'),
    ('操作按钮', '加入编译队列 / 返回'),
]
add_table(['区域', '内容'], result_items)

# ── 3.5 编译队列 ──
doc.add_heading('3.5 编译队列', level=2)
doc.add_paragraph('管理多个客户的编译任务，支持排队、拖拽排序、命令插入。')

doc.add_heading('卡片类型', level=3)
card_types = [
    ('编译卡片', '白色背景', '客户名、项目路径、状态徽章、操作按钮'),
    ('命令卡片', '黄色背景', '命令文字（如 ctvbuild clean）、状态、移除按钮'),
]
add_table(['类型', '样式', '内容'], card_types)

doc.add_heading('状态流转', level=3)
states = [
    ('待编译 (PENDING)', '蓝色', '等待开始'),
    ('编译中 (BUILDING)', '橙色', '正在执行'),
    ('成功 (SUCCEEDED)', '绿色', '编译完成，可上传'),
    ('失败 (FAILED)', '红色', '编译失败，可重试'),
    ('已取消 (CANCELLED)', '灰色', '用户取消'),
]
add_table(['状态', '颜色', '说明'], states)

doc.add_heading('操作功能', level=3)
ops = [
    ('拖拽排序', '拖动卡片左侧手柄调整编译顺序'),
    ('插入 Clean', '在队列末尾插入 ctvbuild clean 命令卡片'),
    ('全部开始编译', '从第一个待编译项开始，依次串行执行'),
    ('自动连编', '成功后自动开始下一个，失败则停止'),
    ('上传到网盘', '编译成功卡片可直接上传产物到网盘'),
    ('清除已完成', '移除所有成功/失败/已取消的卡片'),
]
add_table(['功能', '说明'], ops)

# ── 3.6 编译构建 ──
doc.add_heading('3.6 编译构建', level=2)
doc.add_paragraph('远程编译的实时控制台，通过 tmux 会话执行编译命令。')

build_features = [
    ('编译命令', '自动根据项目/板型/区域/客户/订单生成 ctvbuild 命令，可手动修改'),
    ('实时终端', '黑底终端区域，实时显示编译输出，支持 ANSI 颜色'),
    ('命令输入', '可在终端区域直接输入命令，有编译任务时发到 tmux，否则走 SSH'),
    ('tmux 管理', '查看会话列表、杀掉全部会话'),
    ('进度浮层', '编译启动时显示半透明遮罩 + 转圈进度条'),
    ('自动打包', '编译成功后自动执行 ctvbuild usb 打包'),
    ('选单应答', '编译过程中出现选择菜单时自动按回车确认'),
]
add_table(['功能', '说明'], build_features)

# ── 3.7 规则管理 ──
doc.add_heading('3.7 规则管理', level=2)
doc.add_paragraph('管理所有内置修改规则，支持查看、编辑触发关键词。')

rule_features = [
    ('规则卡片', '170×90 小卡片，按文件分组，Flow 布局自适应排列'),
    ('详情浮层', '点击卡片弹出半透明遮罩，显示规则说明、触发关键词、参数'),
    ('关键词编辑', '内置规则的触发关键词可在浮层中直接编辑保存'),
    ('匹配模式', 'open/close 类规则可编辑动作前缀（打开/开启/启用…）'),
    ('自定义规则', '支持创建自定义规则副本并修改参数'),
]
add_table(['功能', '说明'], rule_features)

# ── 3.8 网盘上传 ──
doc.add_heading('3.8 网盘上传', level=2)
doc.add_paragraph('将编译产物上传到公司 WebDAV 网盘，生成分享链接。')

upload_features = [
    ('网盘设置', '账号/密码/目标文件夹，密码加密存储（keyring）'),
    ('上传队列', '支持多个文件排队上传，逐个浏览服务器添加'),
    ('浏览服务器', '远程文件浏览器，按时间排序，项目快捷按钮'),
    ('批量上传', '全部上传，进度条显示 1/N，逐个显示结果'),
    ('结果卡片', '每个文件独立卡片，显示软件名称/路径，可一键复制'),
    ('编译队列联动', '编译成功卡片可直接加入上传队列，自动跳转上传页'),
]
add_table(['功能', '说明'], upload_features)

# ── 3.9 代码提交 ──
doc.add_heading('3.9 代码提交', level=2)
doc.add_paragraph('将客户目录的修改推送到 Gerrit 代码审核服务器。')

gerrit_features = [
    ('项目选择', '下拉选择项目目录，一键同步最新代码（repo sync）'),
    ('改动文件', 'git status 结果，显示文件状态（M/A/D），可勾选'),
    ('提交信息', '目标分支选择（master/main）+ commit message 输入'),
    ('一键推送', 'git add + commit + push to refs/for/xxx'),
    ('执行日志', '终端风格日志区域，实时显示操作结果'),
]
add_table(['功能', '说明'], gerrit_features)

doc.add_page_break()

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 4. 界面设计要求
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
doc.add_heading('4. 界面设计要求', level=1)

doc.add_heading('4.1 整体布局', level=2)
layout_items = [
    ('布局结构', '左右分栏：左侧导航栏（200px）+ 右侧内容区'),
    ('主背景色', '#e0e3ed（蓝灰色）'),
    ('导航栏背景', '#f5f6f9'),
    ('卡片样式', '白色背景、#e5e5e5 圆角边框（12px）、内边距 16px'),
    ('滚动', '内容区支持滚动，编译构建等页面内部滚动'),
]
add_table(['项目', '规范'], layout_items)

doc.add_heading('4.2 导航栏', level=2)
nav_items = [
    ('1', '自动模式', '🤖', 'AI 分析需求'),
    ('2', '手动模式', '🔧', '可视化编辑配置'),
    ('3', '审核面板', '📋', '审核分析结果'),
    ('4', '执行结果', '📊', '查看修改详情'),
    ('5', '编译队列', '📋', '管理编译任务'),
    ('6', '编译构建', '🔨', '远程编译终端'),
    ('7', '规则管理', '⚙️', '配置修改规则'),
    ('8', '网盘上传', '☁️', '上传编译产物'),
    ('9', '代码提交', '🔀', '推送到 Gerrit'),
]
add_table(['序号', '页面', '图标', '说明'], nav_items)

doc.add_heading('4.3 顶部目录设置栏', level=2)
doc.add_paragraph(
    '除编译队列、编译构建、规则管理、网盘上传、代码提交页面外，'
    '所有页面顶部显示目录设置栏，包含：项目、板型、区域、客户（下拉框），'
    '客户目录（下拉框），目标目录（输入框 + 复制按钮），搜索框（关键词搜索 + 查找/清除按钮）。'
)

doc.add_heading('4.4 交互规范', level=2)
interaction = [
    ('按钮风格', '主按钮：#e0e0e0 背景、#1a1a1a 文字、8px 圆角；小按钮：透明背景、#888 文字、#dcdcdc 边框'),
    ('输入框', '透明背景、#dcdcdc 边框、6px 圆角，聚焦时边框变 #333'),
    ('进度反馈', '耗时操作显示进度浮层（半透明遮罩 + 转圈进度条）'),
    ('错误提示', '红色文字（#ff3b30），支持自动消失'),
    ('成功提示', '绿色文字（#34c759）'),
]
add_table(['项目', '规范'], interaction)

doc.add_heading('4.5 界面布局图', level=2)
doc.add_paragraph('详见附件：docs/layout.html（浏览器打开可查看全部 9 个页面的布局模拟图）。')

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 5. 非功能需求
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
doc.add_heading('5. 非功能需求', level=1)

nfr = [
    ('性能', '页面切换 < 500ms；配置读取 < 3s（SSH）；编译启动 < 5s'),
    ('可靠性', 'SSH 断连自动提示；编译崩溃不丢队列数据（原子写入）'),
    ('安全性', 'WebDAV 密码使用系统 keyring 加密存储；SSH 密钥认证'),
    ('兼容性', '支持 macOS（主要）、Windows；远程服务器为 Linux'),
    ('编码兼容', '配置文件支持 UTF-8、GBK、GB2312、Latin-1 多编码自动识别'),
    ('日志系统', '文件日志（DEBUG 级别）+ 控制台日志（INFO 级别）'),
    ('数据持久化', '编译队列 JSON 文件持久化，单条损坏不影响其他'),
]
add_table(['需求类别', '具体要求'], nfr)

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 6. 技术约束
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
doc.add_heading('6. 技术约束', level=1)

constraints = [
    ('客户端', 'Python 3.13+ / PySide6 / paramiko'),
    ('远程服务器', 'Linux（Ubuntu），SSH 可达，支持 tmux'),
    ('编译工具', 'ctvbuild（内部工具），通过 tmux 会话执行'),
    ('代码管理', 'repo + Gerrit（Android 标准工作流）'),
    ('网盘', 'WebDAV 协议'),
    ('AI 分析', '本地规则引擎 + 可选 LLM 增强'),
    ('配置文件', 'feature_mapping.json 存储规则关键词和匹配模式'),
    ('部署', '单机桌面应用，无需服务端'),
]
add_table(['约束项', '说明'], constraints)

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 7. 版本规划
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
doc.add_heading('7. 版本规划', level=1)

versions = [
    ('v1.0（当前）', '自动/手动模式、审核、编译队列、网盘上传、规则管理、代码提交 UI'),
    ('近期优化', '代码提交 Gerrit 接入、查找最新 zip 优化、编译文件缓存'),
    ('远期规划', '平台化重构（插件架构）、AI 智能分析（LLM）、Web 版本'),
]
add_table(['阶段', '主要内容'], versions)

# ── 保存 ──
out = Path(__file__).parent / '软件输出自动化_产品需求文档.docx'
doc.save(str(out))
print(f'✅ 已生成: {out}')
